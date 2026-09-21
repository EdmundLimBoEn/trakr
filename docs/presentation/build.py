"""Build editable PowerPoint and matching PDF. See README.md for dependencies."""
from pathlib import Path
import textwrap
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

OUT = Path(__file__).resolve().parent
W,H = 960,540
BG='111827'; PANEL='1E293B'; WHITE='F5F3ED'; MUTED='BDC6D5'; LIME='D4F778'; PURPLE='C3B5FD'
fontdir=Path('/usr/share/fonts/truetype/dejavu')
for name,file in [('Deck','DejaVuSans.ttf'),('DeckBold','DejaVuSans-Bold.ttf')]:
    pdfmetrics.registerFont(TTFont(name,str(fontdir/file)))
r=Presentation(); r.slide_width=Inches(13.333333); r.slide_height=Inches(7.5)
r.core_properties.title='Trakr | Equipment accountability, one tap at a time'
r.core_properties.author='Edmund Lim'
c=canvas.Canvas(str(OUT/'Trakr-pitch.pdf'),pagesize=(W,H)); c.setTitle(r.core_properties.title); c.setAuthor('Edmund Lim')
notes=[]; slides=[]; current=None

def rect(x,y,w,h,color,radius=0):
    c.setFillColor('#'+color); c.setStrokeColor('#'+color)
    c.roundRect(x,H-y-h,w,h,radius,fill=1,stroke=0)
    sh=current.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor.from_string(color); sh.line.fill.background()
    if radius: sh.adjustments[0]=0.08

def text(x,y,w,content,size=20,color=WHITE,bold=False):
    font='DeckBold' if bold else 'Deck'; lines=[]
    for para in content.split('\n'):
        line=''
        for word in para.split():
            test=(line+' '+word).strip()
            if pdfmetrics.stringWidth(test,font,size)>w and line: lines.append(line); line=word
            else: line=test
        lines.append(line)
    height=len(lines)*size*1.30+6
    assert y+height <= H+2, (content,y,height)
    c.setFillColor('#'+color); c.setFont(font,size)
    for i,line in enumerate(lines): c.drawString(x,H-y-size-i*size*1.30,line)
    box=current.shapes.add_textbox(Pt(x),Pt(y),Pt(w+4),Pt(height))
    tf=box.text_frame; tf.word_wrap=False
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line
        p.font.name='DejaVu Sans'; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=RGBColor.from_string(color)
        p.space_before=Pt(0); p.space_after=Pt(0); p.line_spacing=1.30
    return height

def begin(section,title,sub,seconds,source,note):
    global current
    current=r.slides.add_slide(r.slide_layouts[6]); n=len(r.slides)
    rect(0,0,W,H,BG); rect(40,34,26,5,LIME)
    text(78,25,780,section.upper(),11,LIME,True)
    title_size = min(34, 870 / pdfmetrics.stringWidth(title, 'DeckBold', 1))
    text(40,66,875,title,title_size,WHITE,True)
    text(40,123,870,sub,15,MUTED)
    rect(40,501,880,1,'374151')
    text(40,513,805,source,8,MUTED)
    text(875,510,60,f'{n:02d}',12,LIME,True)
    current.notes_slide.notes_text_frame.text=f'{seconds}\n\n{note}\n\nEvidence: {source}'
    notes.append(f'## {n:02d}. {title}\n\n**Timing: {seconds}**\n\n{note}\n\nEvidence: {source}\n')

def end(): c.showPage()
def cards(items,y=191,h=235):
    gap=18; width=(880-gap*(len(items)-1))/len(items)
    for i,(head,body) in enumerate(items):
        x=40+i*(width+gap); rect(x,y,width,h,PANEL,12)
        text(x+20,y+18,width-40,f'0{i+1}',12,LIME,True)
        text(x+20,y+51,width-40,head,23,WHITE,True)
        text(x+20,y+(98 if h<200 else 119),width-40,body,15 if h<200 else 16,MUTED)

def banner(s): text(42,455,876,s,17,LIME,True)
def flow(items,y=210,h=155):
    gap=28; width=(880-gap*(len(items)-1))/len(items)
    for i,(head,body) in enumerate(items):
        x=40+i*(width+gap); rect(x,y,width,h,PANEL,10)
        text(x+15,y+18,width-30,head,19,LIME,True)
        text(x+15,y+61,width-30,body,15,WHITE)
        if i<len(items)-1: text(x+width+5,y+55,24,'→',20,PURPLE,True)

begin('01 / The idea','Equipment accountability. One tap at a time.','TRAKR   /   Native iOS + Android   /   School equipment checkout', '0:00–0:20 · 20 seconds','Basis: README.md • native MVP; pilot outcomes not yet measured',
'''A camera leaves the equipment room. Later, a teacher needs to know who has it, what condition it was in, and whether it came back. Trakr connects those moments through an NFC tap, a school account, and a shared record. This is our native mobile MVP and the pilot we propose next.''')
text(45,208,555,'Know who borrowed it.\nKnow what came back.',38,WHITE,True)
text(45,337,505,'A complete handover record for shared school equipment.',21,MUTED)
rect(665,191,250,228,PANEL,18)
text(691,214,205,'TAP → RECORD',19,LIME,True)
text(691,264,200,'ITEM\nBORROWER\nCONDITION\nRETURN',20,WHITE,True)
banner('5-minute pitch  /  3-minute live demonstration'); end()

begin('02 / Problem','The handover is where context gets lost.','Illustrative school scenario: a student borrows a camera and tripod.', '0:20–0:55 · 35 seconds','Problem hypothesis to validate through teacher interviews and baseline observation',
'''Imagine a student borrowing a camera and tripod before a project. A paper log, spreadsheet or chat can capture a name, but keeping the physical handover, condition report and return together takes discipline. Missing context means follow-up work: who should the teacher ask, was the fault already there, and is the item actually back? Our hypothesis is that recording these facts at the handover reduces chasing and uncertainty. We still need to measure the scale of that problem with teachers.''')
cards([('Who has it?','Borrower details can be incomplete or scattered.'),('Was it damaged?','Condition needs to be recorded when custody changes.'),('Did it return?','A return needs a clear closing action and history.')])
banner('Opportunity: make accountability part of the handover itself.'); end()

begin('03 / Users & value','One workflow. Three people who benefit.','Start with one equipment room and one teacher champion.', '0:55–1:30 · 35 seconds','Implemented workflows: README.md • Trakr/Features • dashboard/README.md',
'''Students need a quick way to borrow several items and see their own receipt. Teachers need to enroll equipment, see active claimants, process returns and resolve reported issues. An equipment lead needs oversight across inventory and outstanding activity. Trakr connects those needs with student and teacher mobile workflows plus an operations dashboard. The first target is a single school equipment room, where one teacher can help us observe the whole process before we attempt a broader rollout.''')
cards([('Student','Scan several items. Acknowledge condition. Keep a receipt.'),('Teacher','Enroll tags. View claims. Confirm returns. Resolve issues.'),('Equipment lead','Review inventory, overdue claims and operational exceptions.')])
banner('User value: less ambiguity for students; clearer follow-up for staff.'); end()

begin('04 / Product journey','Tap. Check. Confirm. Return.','Example: a camera with a loose strap and a tripod with no issues.', '1:30–2:10 · 40 seconds','Sources: CheckoutView.swift • ReturnView.swift • Domain.swift',
'''The student signs in with an approved school account and scans both tags. Each item enters a staging list, so accidental scans can be removed. Before confirmation, the student acknowledges condition and adds details for any issue, such as a loose camera strap. Confirmation creates the checkout record and receipt. Later, a teacher scans the returned equipment, reviews active claims and confirms the return. The issue has its own open, acknowledged and resolved lifecycle, so returning an item does not automatically mean its fault has been fixed.''')
flow([('01  Tap','Resolve the NFC tag to an enrolled item.'),('02  Check','Review the list and report condition.'),('03  Confirm','Save claims and show a receipt.'),('04  Return','Teacher closes active claims.')])
banner('Condition reporting and return confirmation are separate responsibilities.'); end()

begin('05 / Data flow','A tap identifies the item. Rules govern the write.','Mobile clients write directly to Firestore using atomic batches.', '2:10–2:50 · 40 seconds','Sources: NFCService.swift • FirebaseGateway.swift • firebase/firestore.rules',
'''Here is the data flow. The NFC tag contains a Trakr identifier, which the phone resolves to an equipment record. Google sign-in provides identity; Firebase Security Rules use verified school email claims to enforce access. On confirmation, the mobile client submits a batch containing the checkout record, individual claims and any issues. Those writes commit together or fail together. A teacher later submits the return batch. The web operations dashboard uses a separate privileged Worker path, with its own authentication and audit events. NFC identifies equipment; it is not proof against tag copying.''')
flow([('NFC tag','tr: identifier\nNo borrower data'),('Native app','Resolve item\nStage + validate'),('Security Rules','Identity + role\nOwnership + state'),('Firestore','Batch + claims\nOptional issues')],205,154)
rect(40,379,880,61,'27334B',8)
text(57,391,845,'Google sign-in → verified identity     |     Teacher return → claim status update',16,WHITE)
banner('Atomic commit: all related writes succeed together, or none do.'); end()

begin('06 / Engineering','Designed for the awkward moments, too.','The quality of the system shows up when a workflow goes wrong.', '2:50–3:25 · 35 seconds','Evidence: docs/mvp-verification.md • TrakrTests • firebase/functions/test',
'''Three engineering decisions matter. First, access is enforced at the database boundary, not just by hiding buttons; students should only see their own claims and issues. Second, the app preserves staged items when offline and blocks confirmation instead of pretending a checkout has succeeded. Third, equipment identity survives tag replacement, preserving claim history. The repository includes tests for these behaviours and repeated operations. Physical NFC, school sign-in and cross-device behaviour still need device validation. We distinguish implemented safeguards from a proven production deployment.''')
cards([('Role boundaries','Verified school identity and record ownership checks.'),('Honest failure','Keep staged items offline; require a connection to confirm.'),('Durable history','Replace a tag without replacing the equipment record.')])
banner('Repository tests exist. Hardware and school-environment proof still matter.'); end()

begin('07 / Validation','Measure the handover, not just the download.','Proposed pilot: one equipment room, four weeks, a teacher sponsor.', '3:25–4:00 · 35 seconds','Proposed experiment • targets are acceptance criteria, not observed results',
'''We propose a four-week pilot. First, observe the existing process and measure its baseline. Then enroll a small inventory and use Trakr for real handovers. Compare median checkout time and weekly teacher follow-up time, while checking record completeness and reliability. Our proposed targets are a thirty percent reduction in checkout time and ninety-five percent complete handover records. These are targets, not results. We also need zero unauthorized cross-student reads in our test set. If the workflow adds friction or staff do not keep using it, we revise before scaling.''')
cards([('Week 1','Baseline timing, teacher interviews and tag enrollment.'),('Weeks 2–3','Observe live handovers, failures and support requests.'),('Week 4','Compare results and decide whether to expand.')])
banner('Targets: ≥30% faster median checkout • ≥95% complete records'); end()

begin('08 / Business case','Start narrow. Earn the right to expand.','Business hypothesis: a school subscription with supported onboarding.', '4:00–4:35 · 35 seconds','Commercial proposal • no validated pricing, paying customers or market-size claim',
'''The proposed buyer is the school or equipment-owning department, while a teacher is the champion. A subscription could fund hosting, support and continued development, with onboarding covering tags and inventory setup. We would first win one room, then expand within the school after showing repeat use and time saved. Today the differentiation is the complete handover workflow, not a proprietary NFC technology. The investor milestone is repeatable adoption and willingness to pay, followed by the multi-school architecture and onboarding process needed to scale. Revenue and pricing remain unvalidated.''')
cards([('Buyer','School or department, supported by a teacher champion.'),('Revenue hypothesis','Annual subscription; optional tag and setup support.'),('Expansion','One room → more departments → additional schools.')])
banner('Next proof: repeat usage, willingness to pay and support cost per school.'); end()

begin('09 / The ask','Give us one room to prove the value.','A focused pilot turns a working MVP into evidence.', '4:35–5:00 · 25 seconds','Ask: proposed pilot partnership • demo begins at 5:00',
'''Our ask is one equipment room, one teacher sponsor and permission to run a measured pilot. We will return a baseline comparison, reliability findings and an adoption recommendation. Trakr already connects checkout, condition and return in a working software flow. The next question is whether it reliably makes school handovers easier. Let us now show that loop in three minutes.''')
text(45,204,855,'1 room. 1 teacher. 4 weeks.',43,LIME,True)
text(45,291,820,'Deliverable: measured time savings, record quality and a clear go / no-go decision.',25,WHITE)
text(45,407,820,'NEXT  /  Watch one student checkout become a teacher-confirmed return.',18,PURPLE,True); end()

begin('Demo / Keep this slide on screen','Three minutes. One complete handover.','Show the record changing, not every menu in the app.', '5:00–8:00 · 180 seconds','Use prepared demo inventory; local demo proves UI flow, not cross-device sync',
'''0:00–0:20: Introduce the prepared student account and two enrolled items. 0:20–1:00: Scan the camera and tripod; show the staging list. 1:00–1:35: Report a loose strap on the camera, acknowledge condition, confirm and show the receipt. 1:35–2:20: Switch to the prepared teacher session, scan the same items, show the claimant and confirm return. 2:20–2:45: Show the returned record and acknowledge the issue, explaining that repair is a separate action. 2:45–3:00: Close with: one handover, one shared record, a clear next action. Rehearse account switching. Use a dedicated test environment. If physical NFC fails, switch promptly to the prepared local demo and explicitly identify it as simulated. Independent local demos do not demonstrate cross-device synchronization.''')
flow([('0:00–1:00','Student + two items\nScan into staging'),('1:00–1:35','Report condition\nConfirm + receipt'),('1:35–2:20','Teacher reviews\nConfirm return'),('2:20–3:00','History + issue\nExplain next action')],209,169)
banner('Success: the audience sees who borrowed it, its condition and its return.'); end()

begin('Appendix A / Technical depth','Identity stays stable when tags change.','Simplified data relationships; names follow the Firebase implementation.', 'Q&A only','Sources: Domain.swift • FirebaseGateway.swift • firebase/firestore.rules',
'''Equipment is the stable identity. A tag points to equipment; tag replacement changes that pointer while existing claims retain their tag-at-checkout reference. A checkout batch groups claims, and each claim belongs to a student and equipment item. An issue references a claim and equipment. A return batch groups claim closures with a teacher identity. Serial guard records support uniqueness. Multiple active claims per equipment item are currently allowed, so do not describe this as an exclusive booking system.''')
flow([('tags','equipmentId\nstatus'),('equipment','name + serial\nstable identity'),('claims','student + item\ncondition + status'),('issues','claim + item\nopen → resolved')],194,151)
text(50,372,860,'checkoutBatches → claims ← returnBatches',24,PURPLE,True)
text(50,424,850,'Tag replacement preserves history. Multiple active claims are currently allowed.',17,MUTED); end()

begin('Appendix B / Trust boundaries','Two access paths. Two enforcement points.','Mobile database access and privileged operations access must both be secured.', 'Q&A only','Sources: firebase/firestore.rules • dashboard/src/worker/auth.ts • ops.ts',
'''Mobile: Google Authentication supplies verified identity, and Firestore Security Rules govern reads and writes. Staff domains map to teacher access, student subdomains map to student access. Operations: the browser signs in, the Worker validates the identity and session, then accesses Firestore using a server-held service account. That privileged path does not rely on client Security Rules; its own authorization and validation are essential. Ops mutations generate audit events. The shared-secret emergency route is a separate operational risk. Multi-school deployment requires redesigned tenant boundaries; changing branding alone is insufficient.''')
flow([('Student / teacher','Google identity\nNative client'),('Firestore Rules','Role + ownership\nSchema + state'),('Firestore','Equipment\nClaims + issues')],183,131)
flow([('Staff browser','Google sign-in\nSession cookie'),('Cloudflare Worker','Authorize + validate\nServer-held credential'),('Firestore','Privileged ops\nMutation audit')],338,131); end()

begin('Appendix C / Honest scope','An MVP with clear limits.', 'These constraints define the next engineering work.', 'Q&A only','Sources: docs/production-integration.md • docs/mvp-verification.md • README.md',
'''Overdue status is calculated in-app and alerts are in-app, not guaranteed push delivery. An internet connection is required to confirm. NFC read/write still needs physical device proof, and school Google Workspace access requires validation. The current role model is school-specific. Multiple claims are allowed, and the return workflow closes the active claims it loads; concurrency deserves pilot testing. Demo accounts have deliberately relaxed write rules and share collections with school accounts in the documented setup. Use a dedicated demo environment and disable demo access before public launch. Do not claim GPS location, theft prevention, exclusive reservations or commercial traction.''')
cards([('Connectivity','Confirmation requires internet. Overdue checks run in-app.'),('Identity & custody','School-specific roles. NFC tags are identifiers, not anti-theft devices.'),('Launch readiness','Validate hardware, school OAuth, concurrency and demo isolation.')],h=249)
banner('Roadmap: push delivery, tenant isolation and stronger operational validation.'); end()

begin('Appendix D / Evaluation','Make each claim testable.', 'A proposed assessment map, since no course rubric was supplied.', 'Q&A only','Evidence map: docs/mvp-verification.md and source references in speaker notes',
'''This is a suggested assessment map, not an official grading rubric. Problem definition is supported by the concrete borrowing scenario, then strengthened with interviews. Design understanding is shown through the user journey, architecture and data relationships. Implementation is demonstrated by the complete checkout-to-return loop. Quality is supported by repository test coverage and explicit device validation gaps. Critical evaluation means explaining the tradeoffs and measuring against a baseline, rather than claiming the project is production-ready. Source traceability and honest status labels let a teacher separate implementation evidence from proposed outcomes.''')
rows=[('Problem understanding','User scenario + teacher interviews to collect'),('Design & implementation','Journey + data flow + live end-to-end demo'),('Testing & reliability','Rules/store/UI tests in repository; device checks pending'),('Critical evaluation','Constraints, alternatives and pilot acceptance criteria'),('Communication','Five-minute argument + three-minute demonstration')]
for i,(a,b) in enumerate(rows):
    y=181+i*58; rect(40,y,880,49,PANEL,5); text(54,y+10,272,a,16,LIME,True); text(341,y+10,564,b,15,WHITE)
end()

begin('Appendix E / Alternatives','Why this workflow instead of another log?', 'Conceptual comparison; no vendor benchmarking has been performed.', 'Q&A only','Analysis of workflow options, not verified competitor capability claims',
'''Paper and spreadsheets are familiar and can be sufficient for low-volume borrowing. QR-based workflows are also a credible option; NFC must earn its place through device suitability and observed convenience, rather than an unsupported speed claim. A broad asset platform may be appropriate for institutions that need procurement, depreciation or integrations; those requirements are outside this MVP. Trakr focuses on linking the physical handover, identity, condition, receipt and teacher-confirmed return. We should choose the simplest solution that solves the observed problem. A defensible business would require trusted deployment and repeatable onboarding, not just an NFC tag.''')
cards([('Paper / sheets','Low setup burden. Handover completeness depends on process discipline.'),('QR / NFC forms','A credible lightweight option. Compare scan reliability and completion.'),('Trakr focus','A connected checkout, condition, receipt and teacher-return workflow.')],h=249)
banner('Differentiation hypothesis: workflow completion, not the tag technology.'); end()

begin('Appendix F / Pilot scorecard','Set the decision rule before the pilot.', 'Record baseline and pilot samples; report sample size and failures.', 'Q&A only','All thresholds below are proposed pilot targets, not achieved measurements',
'''Time the same defined task before and during the pilot: beginning the handover through a confirmed record, separating student task time from teacher follow-up. Use comparable item counts and report the number of observations, median and slow cases. Define complete records as having borrower, equipment, checkout time, condition and a teacher-confirmed return when the item is physically returned. Review adoption weekly and interview staff about workarounds. Security tests and failed commits are mandatory checks, not substitutes for usability results. Agree thresholds with the teacher sponsor before the experiment.''')
rows=[('Median checkout time','Target: ≥30% below the observed baseline'),('Record completeness','Target: ≥95% of eligible handovers complete'),('Student data isolation','Target: 0 unauthorized reads in the test set'),('Workflow reliability','Track failed commits, duplicates and recovery'),('Adoption & follow-up','Measure weekly use and staff minutes spent chasing')]
for i,(a,b) in enumerate(rows):
    y=181+i*58; rect(40,y,880,49,PANEL,5); text(54,y+10,280,a,16,LIME,True); text(351,y+10,550,b,15,WHITE)
end()

begin('Appendix G / Investor questions','Prove willingness to pay before scaling.', 'Early business exploration; there is no validated revenue forecast.', 'Q&A only','Illustrative arithmetic only • source license: LICENSE (CC BY-NC 4.0)',
'''Illustration only: if a school saves two staff hours each week, values time at thirty Singapore dollars per hour and uses the system for forty weeks, gross time value is two thousand four hundred dollars per school year. This is not cash savings or validated willingness to pay. Net value must account for onboarding, tags, support and hosting; the subscription must leave the buyer better off. Reachable market should be built from named qualified schools and realistic adoption, not a fabricated global TAM. Before commercial expansion, clarify rights for a commercial offering given the repository's noncommercial license, and validate procurement and data requirements. Multi-tenant isolation is future engineering work.''')
text(45,190,870,'2 hours × S$30 × 40 weeks = S$2,400',31,LIME,True)
text(45,248,850,'Illustrative annual time value per school — not measured savings or revenue.',17,MUTED)
cards([('Discover','Who owns the budget? What evidence unlocks a purchase?'),('Cost','Tags, setup, hosting, support and acquisition all count.'),('Gate expansion','Paid conversion, retention, tenant isolation and commercial rights.')],y=305,h=174)
end()

begin('Appendix H / Evidence & next steps','A pitch grounded in the repository.', 'Implementation evidence is not the same as field validation.', 'Q&A only','Repository snapshot: 6467a41 • presentation prepared 21 September 2026',
'''The presentation is grounded in README.md, the domain models, FirebaseGateway, NFCService, Firestore rules and operations dashboard documentation. The verification matrix records repository test coverage and explicitly calls out hardware and school-infrastructure gaps. We did not rerun mobile or backend suites to produce this deck, and we do not claim a new test pass. The next evidence to collect is teacher interviews, baseline handover timings, physical-device results, pilot usage and buyer feedback. Proposed pricing, targets and commercial paths must be updated when real evidence arrives.''')
rows=[('Product scope','README.md · docs/production-integration.md'),('Data & mutation flow','Trakr/Models/Domain.swift · Services/FirebaseGateway.swift'),('Security & operations','firebase/firestore.rules · dashboard/README.md'),('Verification coverage','docs/mvp-verification.md · TrakrTests · firebase/functions/test'),('Next evidence','Teacher interviews · device results · pilot data · buyer feedback')]
for i,(a,b) in enumerate(rows):
    y=181+i*58; rect(40,y,880,49,PANEL,5); text(54,y+10,240,a,16,LIME,True); text(304,y+10,597,b,14,WHITE)
end()

r.save(OUT/'Trakr-pitch.pptx'); c.save()
(OUT/'speaker-notes.md').write_text('# Trakr presentation script\n\nSlides 1–9: five minutes. Slide 10: three-minute demo. Slides 11–18: Q&A only.\n\n'+'\n'.join(notes))
print(f'Built {len(r.slides)} slides with speaker notes.')
